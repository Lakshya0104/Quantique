import json, copy
from pptx import Presentation
from pptx.util import Inches as I, Pt, Emu
from pptx.dml.color import RGBColor as C
from pptx.enum.shapes import MSO_SHAPE as S, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN as A, MSO_ANCHOR as V
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE as CT, XL_LEGEND_POSITION as LP, XL_LABEL_POSITION as LBL
from pptx.oxml.ns import qn
D=json.load(open('data.json'))
PINK=C(0xF2,0x37,0xA6); INK=C(0x12,0x12,0x16); GREY=C(0x5E,0x5E,0x6A); LG=C(0xF4,0xF4,0xF6); TINT=C(0xFD,0xEB,0xF5); MID=C(0xA8,0xA8,0xB2); WHITE=C(255,255,255); DEEP=C(0x8A,0x14,0x5E)
SERIF='Cambria'; SANS='Calibri'; MONO='Consolas'
p=Presentation('/tmp/claude-0/t.pptx'); SL=list(p.slides)

def title(s,t):
    for sh in s.shapes:
        if sh.name=='TextBox 3':
            para=sh.text_frame.paragraphs[0]; rs=para.runs
            rs[0].text=t
            for r in rs[1:]: r._r.getparent().remove(r._r)
            f=para.runs[0].font; f.name=SERIF; f.size=Pt(26)
def tb(s,x,y,w,h,runs,size=12,color=INK,font=SANS,bold=False,align=A.LEFT,anchor=V.TOP,spc=None,ls=None,italic=False,after=0):
    b=s.shapes.add_textbox(I(x),I(y),I(w),I(h)); tf=b.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0; tf.vertical_anchor=anchor
    if isinstance(runs,str): runs=[[(runs,{})]]
    elif runs and isinstance(runs[0],tuple): runs=[runs]
    for i,para in enumerate(runs):
        pp=tf.paragraphs[0] if i==0 else tf.add_paragraph(); pp.alignment=align
        if ls: pp.line_spacing=ls
        pp.space_after=Pt(after)
        if isinstance(para,str): para=[(para,{})]
        for txt,o in para:
            r=pp.add_run(); r.text=txt; f=r.font
            f.name=o.get('font',font); f.size=Pt(o.get('size',size)); f.bold=o.get('bold',bold); f.italic=o.get('italic',italic)
            f.color.rgb=o.get('color',color)
            if o.get('link'): r.hyperlink.address=o['link']
            sp=o.get('spc',spc)
            if sp: r._r.get_or_add_rPr().set('spc',str(sp))
    return b
def box(s,x,y,w,h,fill=None,line=None,shape=S.RECTANGLE,lw=1,radius=None):
    r=s.shapes.add_shape(shape,I(x),I(y),I(w),I(h))
    if fill is None: r.fill.background()
    else: r.fill.solid(); r.fill.fore_color.rgb=fill
    if line is None: r.line.fill.background()
    else: r.line.color.rgb=line; r.line.width=Pt(lw)
    r.shadow.inherit=False
    if radius is not None and shape==S.ROUNDED_RECTANGLE: r.adjustments[0]=radius
    return r
def arrow(s,x1,y1,x2,y2,color=INK,w=1.25):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,I(x1),I(y1),I(x2),I(y2)); c.line.color.rgb=color; c.line.width=Pt(w)
    ln=c.line._get_or_add_ln(); t=ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}); ln.append(t); return c
def eyebrow(s,x,y,t,color=PINK,w=4,align=A.LEFT): return tb(s,x,y,w,0.22,t.upper(),size=9,color=color,font=MONO,bold=True,spc=200,align=align)
def head(s,t): tb(s,1.2,1.38,10.93,0.4,t,size=15,color=GREY,font=SERIF,italic=True,align=A.CENTER)
def circ(s,x,y,d,txt,fill=PINK,fc=WHITE,size=12):
    c=box(s,x,y,d,d,fill=fill,shape=S.OVAL); tf=c.text_frame; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    pp=tf.paragraphs[0]; pp.alignment=A.CENTER; r=pp.add_run(); r.text=txt; r.font.size=Pt(size); r.font.bold=True; r.font.color.rgb=fc; r.font.name=SANS; tf.vertical_anchor=V.MIDDLE; return c
def img(s,path,x,y,w=None,h=None):
    return s.shapes.add_picture(path,I(x),I(y),I(w) if w else None,I(h) if h else None)
def style_chart(ch,size=10):
    ch.font.size=Pt(size); ch.font.name=SANS; ch.font.color.rgb=GREY
def card(s,x,y,w,h,num,ttl,body,fill=LG,bs=11):
    box(s,x,y,w,h,fill=fill)
    eyebrow(s,x+0.18,y+0.15,num,w=w-0.3)
    tb(s,x+0.18,y+0.4,w-0.36,0.35,ttl,size=14,bold=True,font=SERIF)
    tb(s,x+0.18,y+0.78,w-0.36,h-0.85,body,size=bs,color=GREY,ls=1.05)

# ---------- 1 TITLE
s=SL[0]; title(s,'QUANTUM PNT FOR DRONES')
eyebrow(s,0.75,1.75,'UC-085 · Defence & Strategic Applications',w=6)
tb(s,0.75,2.05,6.2,1.0,'VOID-NAV',size=54,bold=True,font=SERIF)
tb(s,0.75,2.95,5.9,1.1,'Navigating without the sky: a quantum-assisted positioning payload that reads the Earth\'s magnetic fingerprint when GNSS is jammed, spoofed or unavailable.',size=15,color=GREY,font=SERIF,ls=1.1)
chips=['Quantum sensing','Grover search','Quantum kernel ML','QAOA planning']
x=0.75
for c in chips:
    w=0.25+len(c)*0.078; b=box(s,x,4.15,w,0.34,fill=TINT,shape=S.ROUNDED_RECTANGLE,radius=0.5)
    tb(s,x,4.15,w,0.34,c,size=10.5,color=DEEP,bold=True,align=A.CENTER,anchor=V.MIDDLE); x+=w+0.12
eyebrow(s,0.75,4.85,'Team',color=GREY); tb(s,0.75,5.08,3,0.45,'Team VOID',size=24,bold=True,font=SERIF)
eyebrow(s,3.4,4.85,'Deliverable',color=GREY); tb(s,3.4,5.12,3.5,0.5,'Working payload prototype + Qiskit pipeline + live dashboard',size=12,color=INK)
img(s,'drone.png',7.15,1.6,w=5.75)
box(s,7.3,5.05,5.4,0.62,fill=TINT,shape=S.ROUNDED_RECTANGLE,radius=0.2)
tb(s,7.5,5.05,1.9,0.62,'VIDEO PITCH',size=10,font=MONO,bold=True,color=PINK,anchor=V.MIDDLE,spc=100)
tb(s,9.35,5.05,3.3,0.62,[('https://youtu.be/BV2KYV-yhL0',{'link':'https://youtu.be/BV2KYV-yhL0','color':DEEP,'bold':True})],size=14,anchor=V.MIDDLE)

# ---------- 2 PROBLEM
s=SL[1]; title(s,'PROBLEM STATEMENT'); head(s,'Without GNSS, a drone\'s inertial position drifts without bound.')
stats=[('10⁻¹⁶ W','GNSS power at the receiver: weaker than background noise, easily denied by jamming, terrain, tunnels or foliage.'),
 ('~260 m','Mean error after 3 min of MEMS-only dead reckoning in our drift model: error grows with t², not t.'),
 ('₹10L+','Cost of navigation-grade INS or cold-atom units: too expensive to fit on every tactical drone.')]
for i,(k,v) in enumerate(stats):
    y=1.95+i*1.12; box(s,0.7,y,4.6,0.98,fill=LG)
    tb(s,0.9,y+0.12,1.7,0.7,k,size=24,bold=True,font=SERIF,color=PINK,anchor=V.MIDDLE)
    tb(s,2.55,y+0.1,2.6,0.8,v,size=10.5,color=GREY,anchor=V.MIDDLE,ls=1.0)
cd=CategoryChartData(); cd.categories=[f'{t}s' for t in D['t']]
cd.add_series('Inertial only (MEMS IMU)',D['ins']); cd.add_series('VOID-NAV (MEMS mag fix)',D['mems'])
gf=s.shapes.add_chart(CT.LINE,I(5.6),I(1.85),I(7.1),I(3.2),cd); ch=gf.chart; style_chart(ch)
ch.has_title=True; ch.chart_title.text_frame.text='Position error after GPS loss (m), simulated'; tf=ch.chart_title.text_frame.paragraphs[0].runs[0].font; tf.size=Pt(11); tf.bold=True; tf.color.rgb=INK; tf.name=SANS
ch.has_legend=True; ch.legend.position=LP.BOTTOM; ch.legend.include_in_layout=False
for sr,col in zip(ch.plots[0].series,[MID,PINK]):
    sr.format.line.color.rgb=col; sr.format.line.width=Pt(2.75); sr.smooth=True
va=ch.value_axis; va.major_gridlines.format.line.color.rgb=C(0xE6,0xE6,0xEA); va.format.line.fill.background(); va.has_major_gridlines=True
ch.category_axis.format.line.color.rgb=MID; ch.category_axis.tick_labels.font.size=Pt(9)
box(s,0.7,5.3,12.0,0.8,fill=TINT)
tb(s,0.95,5.36,11.5,0.68,[('Problem: ',{'bold':True,'color':DEEP}),('give a defence UAV a position it can trust within metres when GNSS disappears, with a passive (non-emitting), low-cost payload that is ready to take a quantum magnetometer, as UC-085 and the iDEX passive quantum positioning challenge call for.',{})],size=12.5,anchor=V.MIDDLE,font=SERIF)

# ---------- 3 SOLUTION
s=SL[2]; title(s,'PROPOSED SOLUTION'); head(s,'VOID-NAV: a passive quantum PNT payload. The drone carries it; the payload does the navigation.')
img(s,'drone.png',0.45,1.95,w=6.4)
steps=[('Map','GPS on: fly or walk the area once and log a 5-D magnetic fingerprint [Bx, By, Bz, |B|, ∇B] per 2 m cell.'),
 ('Sense','GPS off: the magnetometer, IMU, optical flow and barometer track motion; the magnetic field is read as a short window.'),
 ('Match','Quantum engine: Grover\'s search finds candidate cells, the quantum kernel SVM confirms the zone.'),
 ('Correct','A particle filter fuses the match with dead reckoning, resets IMU drift and sends position + confidence to the autopilot.')]
for i,(t,b) in enumerate(steps):
    y=1.95+i*1.03
    circ(s,7.15,y+0.05,0.5,str(i+1),size=14)
    if i<3: c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,I(7.4),I(y+0.58),I(7.4),I(y+1.05)); c.line.color.rgb=PINK; c.line.width=Pt(1.25); c.line.dash_style=4
    tb(s,7.85,y,4.9,0.32,t,size=15,bold=True,font=SERIF)
    tb(s,7.85,y+0.33,4.9,0.65,b,size=11,color=GREY,ls=1.0)

# ---------- 4 INNOVATION
s=SL[3]; title(s,'INNOVATION & UNIQUENESS'); head(s,'Most entries show one quantum algorithm. VOID-NAV runs a complete quantum PNT pipeline.')
cards=[('01 · Full stack','Sense, search, learn, plan','Quantum sensing, Grover\'s search, quantum-kernel ML and QAOA each do a distinct job in one navigation loop.'),
('02 · Richer fingerprint','5-D magnetic gradient map','We match the spatial pattern [Bx, By, Bz, |B|, ∇B], not a single field value, which cuts false matches sharply.'),
('03 · Quantum cross-check','Grover proposes, QSVM decides','In our run Grover returned two look-alike cells (0110, 1011); the QSVM zone label resolved the tie.'),
('04 · Quantum-sensor ready','Drop-in NV-diamond slot','Same data interface as an NV magnetometer; we simulate its Ramsey readout today to quantify the upgrade.'),
('05 · Magnetically aware routes','QAOA path planning','Chooses waypoints through magnetically rich cells so the drone gets more position fixes per minute.'),
('06 · Deployable','Passive, retrofit, low cost','Emits nothing, speaks MAVLink to any ArduPilot/PX4 drone, prototype under ₹8,000 in parts.')]
for i,(n,t,b) in enumerate(cards):
    x=0.7+(i%3)*4.05; y=1.9+(i//3)*1.75
    card(s,x,y,3.85,1.6,n,t,b)

pl=[('SENSE','Ramsey + IPE'),('SEARCH','Grover'),('LEARN','QSVM'),('PLAN','QAOA'),('FUSE','Particle filter')]
for i,(a,b2) in enumerate(pl):
    x=0.7+i*2.46; bx=box(s,x,5.45,2.1,0.62,fill=INK if i<4 else PINK)
    tb(s,x,5.49,2.1,0.25,a,size=9,font=MONO,bold=True,color=C(0xF9,0xA8,0xD8) if i<4 else WHITE,align=A.CENTER,spc=150)
    tb(s,x,5.73,2.1,0.3,b2,size=11.5,bold=True,color=WHITE,align=A.CENTER)
    if i<4: arrow(s,x+2.12,5.76,x+2.44,5.76,color=PINK,w=1.5)
# ---------- 5 QUANTUM APPROACH
s=SL[4]; title(s,'QUANTUM COMPUTING APPROACH'); head(s,'Four quantum building blocks, each mapped to one navigation job.')
panels=[('Q1 · Quantum sensing','Ramsey interferometry + iterative phase estimation','ipe.png','Qubit in superposition via H; field B imprints phase φ = γBτ (γ ≈ 28 GHz/T); phase kickback + IPE reads φ bit by bit. Models an NV-diamond readout.'),
('Q2 · Search','Grover\'s algorithm · 4 qubits, 16 map cells','grover.png','Oracle U_f phase-flips cells whose fingerprint matches; diffuser 2|s⟩⟨s|−I amplifies them. ~√N queries; 2 iterations here.'),
('Q3 · Learning','Quantum-kernel SVM · ZZ feature map','zz.png','x ↦ |φ(x)⟩ in a 16-dim Hilbert space; kernel K(x,x′)=|⟨φ(x)|φ(x′)⟩|² feeds a classical SVM (QSVC) to label the zone.'),
('Q4 · Planning','QAOA · p = 1, route graph','qaoa.png','Cost Hamiltonian H_C = Σ wᵢⱼZᵢZⱼ rewards magnetically rich waypoints; mixer H_M = ΣXᵢ; γ, β tuned by COBYLA.')]
for i,(n,t,im,b) in enumerate(panels):
    x=0.7+(i%2)*6.1; y=1.9+(i//2)*2.15; w=5.9; h=2.0
    box(s,x,y,w,h,fill=LG)
    eyebrow(s,x+0.18,y+0.12,n,w=3)
    tb(s,x+0.18,y+0.36,w-0.36,0.3,t,size=12.5,bold=True,font=SERIF)
    pic=img(s,im,x+0.18,y+0.66,h=0.97)
    mw=I(w-0.36)
    if pic.width>mw:
        r=mw/pic.width; pic.width=mw; pic.height=int(pic.height*r); pic.top=int(I(y+0.66)+(I(0.97)-pic.height)/2)
    tb(s,x+0.18,y+1.62,w-0.36,0.4,b,size=9,color=GREY,ls=1.0)

# ---------- 6 STACK
s=SL[5]; title(s,'QISKIT & TECHNOLOGY STACK'); head(s,'Open-source software on borrowed hardware: built for a student budget.')
layers=[('Quantum','Qiskit 2.x SDK · Qiskit Aer · IBM Runtime SamplerV2 · qiskit-machine-learning (FidelityQuantumKernel, QSVC) · QAOA + COBYLA'),
('Navigation','Python · NumPy / SciPy · FilterPy particle filter + EKF · hard/soft-iron calibration'),
('Flight','ArduPilot SITL · pymavlink · MAVLink VISION_POSITION_ESTIMATE / GPS_INPUT'),
('Firmware','ESP32 · Arduino C++ · I²C sensor drivers · 50 Hz UDP stream'),
('Interface','Streamlit + Plotly live map · circuit + histogram panels · CSV / Parquet logs')]
for i,(k,v) in enumerate(layers):
    y=1.95+i*0.8; box(s,0.7,y,6.55,0.7,fill=TINT if i==0 else LG)
    tb(s,0.9,y,1.4,0.7,k.upper(),size=9.5,bold=True,font=MONO,color=PINK,anchor=V.MIDDLE,spc=100)
    tb(s,2.3,y,4.85,0.7,v,size=10.5,color=INK,anchor=V.MIDDLE,ls=1.0)
rows=[('Part','Role','Source','₹'),('QMC5883L magnetometer','Main field sensor (mast)','Buy','300'),('MPU-9250 IMU','Accel + gyro','Borrow','600'),('BMP280 barometer','Altitude','Borrow','150'),('NEO-6M GPS','Mapping + ground truth','Borrow','450'),('PMW3901 optical flow','Ground motion','Buy (opt.)','2,000'),('ESP32 DevKit','Sensor hub','Borrow','450'),('Quadcopter frame / drone','Platform','Borrow','0'),('Total if bought','','','≈ 3,950')]
t=s.shapes.add_table(len(rows),4,I(7.5),I(1.95),I(5.2),I(3.9)).table
for j,wd in enumerate([2.05,1.65,0.85,0.65]): t.columns[j].width=I(wd)
for i,r in enumerate(rows):
    for j,v in enumerate(r):
        c=t.cell(i,j); c.text=v or ' '; c.margin_left=c.margin_right=I(0.06); c.margin_top=c.margin_bottom=I(0.03)
        c.fill.solid(); c.fill.fore_color.rgb=INK if i==0 else (TINT if i==len(rows)-1 else (WHITE if i%2 else LG))
        f=c.text_frame.paragraphs[0]; f.alignment=A.RIGHT if j==3 else A.LEFT
        rr=f.runs[0].font; rr.size=Pt(9.5); rr.name=SANS; rr.bold=(i==0 or i==len(rows)-1); rr.color.rgb=WHITE if i==0 else INK
tt=t._tbl; tt.getparent().getparent().getparent()
# ---------- 7 ARCHITECTURE
s=SL[6]; title(s,'SYSTEM ARCHITECTURE / WORKFLOW'); head(s,'Sensors in, trusted position out: quantum where it adds value, classical where speed matters.')
# column 1 sensors
eyebrow(s,0.7,1.9,'Sense · 50 Hz'); 
for i,(a,b) in enumerate([('Magnetometer','quantum-ready'),('IMU','accel · gyro'),('Optical flow','ground motion'),('Barometer','altitude'),('GPS','mapping only')]):
    y=2.2+i*0.72; box(s,0.7,y,2.2,0.6,fill=LG); tb(s,0.85,y+0.05,2.0,0.28,a,size=11.5,bold=True); tb(s,0.85,y+0.31,2.0,0.25,b,size=9.5,color=GREY)
arrow(s,2.95,3.95,3.35,3.95)
eyebrow(s,3.4,1.9,'Pre-process')
b=box(s,3.4,2.2,2.1,3.48,fill=WHITE,line=INK,lw=1)
tb(s,3.55,2.35,1.85,3.2,[[('Calibrate',{'bold':True})],[('hard/soft iron, tilt',{'color':GREY,'size':9.5})],[(' ',{})],[('Fingerprint',{'bold':True})],[('[Bx,By,Bz,|B|,∇B]',{'color':GREY,'size':9.5,'font':MONO})],[(' ',{})],[('Dead reckoning',{'bold':True})],[('IMU + flow → Δx, Δy',{'color':GREY,'size':9.5})]],size=11.5)
arrow(s,5.55,3.95,5.95,3.95)
eyebrow(s,6.0,1.9,'Quantum engine · Qiskit')
box(s,6.0,2.2,3.35,3.48,fill=INK)
q=[('Q1 Ramsey / IPE','read field (NV sim)'),('Q2 Grover','candidate cells'),('Q3 QSVM','zone label'),('Q4 QAOA','next waypoints')]
for i,(a,bb) in enumerate(q):
    y=2.35+i*0.82; box(s,6.15,y,3.05,0.7,fill=C(0x26,0x26,0x2E))
    tb(s,6.3,y+0.07,2.8,0.3,a,size=11.5,bold=True,color=WHITE); tb(s,6.3,y+0.37,2.8,0.28,bb,size=9.5,color=C(0xF9,0xA8,0xD8))
arrow(s,9.4,3.95,9.8,3.95)
eyebrow(s,9.85,1.9,'Fuse & act')
box(s,9.85,2.2,2.85,1.6,fill=TINT); tb(s,10.0,2.32,2.6,1.4,[[('Particle filter',{'bold':True,'size':12.5})],[('500 particles; re-weighted only when Grover and QSVM agree',{'color':GREY,'size':10})]],size=11.5)
arrow(s,11.27,3.85,11.27,4.1)
box(s,9.85,4.1,1.37,1.58,fill=LG); tb(s,9.95,4.2,1.2,1.4,[[('Autopilot',{'bold':True})],[('MAVLink position + covariance',{'color':GREY,'size':9.5})]],size=11)
box(s,11.33,4.1,1.37,1.58,fill=LG); tb(s,11.43,4.2,1.2,1.4,[[('Dashboard',{'bold':True})],[('map, error, circuits',{'color':GREY,'size':9.5})]],size=11)
tb(s,0.7,5.82,12,0.3,'Loop: classical fusion runs every 20 ms; quantum matching runs every 1 s on a window of readings, so quantum latency never blocks flight control.',size=10.5,color=GREY,italic=True)

# ---------- 8 FEASIBILITY
s=SL[7]; title(s,'IMPLEMENTATION & FEASIBILITY'); head(s,'Built in 24 hours, with a working fallback at every level.')
eyebrow(s,0.7,1.9,'24-hour build plan')
X0=2.1; W=10.6
for h in range(0,25,4):
    x=X0+W*h/24; tb(s,x-0.3,2.15,0.6,0.22,f'H{h}',size=9,font=MONO,color=GREY,align=A.CENTER)
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,I(x),I(2.4),I(x),I(4.15)); c.line.color.rgb=C(0xE0,0xE0,0xE6); c.line.width=Pt(0.75)
lanes=[('Hardware',[(0,5,'Wire sensors + ESP32 firmware',INK),(5,9,'Calibrate + mapping run',GREY),(16,20,'Mount on drone, test',MID)]),
('Quantum',[(0,6,'Grover + IPE circuits (Aer)',PINK),(6,12,'QSVM on map data',DEEP),(12,16,'QAOA + IBM run',C(0xF7,0x8C,0xCB))]),
('Nav + UI',[(2,10,'Dead reckoning + particle filter',INK),(12,16,'Integrate quantum → filter',GREY),(16,20,'Dashboard',MID)])]
for i,(n,items) in enumerate(lanes):
    y=2.5+i*0.55; tb(s,0.7,y,1.3,0.42,n,size=11,bold=True,anchor=V.MIDDLE)
    for a,b2,t,col in items:
        r=box(s,X0+W*a/24+0.02,y,W*(b2-a)/24-0.04,0.42,fill=col); tb(s,X0+W*a/24+0.08,y,W*(b2-a)/24-0.12,0.42,t,size=9,color=WHITE if col not in (MID,C(0xF7,0x8C,0xCB)) else INK,anchor=V.MIDDLE)
box(s,X0+W*20/24+0.02,2.5,W*4/24-0.04,1.52,fill=TINT); tb(s,X0+W*20/24+0.1,2.5,W*4/24-0.2,1.52,[[('Rehearse',{'bold':True})],[('record backup video',{'size':9,'color':GREY})]],size=11,anchor=V.MIDDLE,color=DEEP)
lv=[('Level 1 · guaranteed','Payload carried along the mapped corridor; live dashboard shows drift vs correction.'),('Level 2 · main demo','Payload mounted on a borrowed drone, carried or tethered; full flight replayed in ArduPilot SITL on our real map.'),('Level 3 · bonus','Short real flight with GPS disabled in software, only if a flight-ready drone and venue permission exist.')]
for i,(a,b2) in enumerate(lv):
    x=0.7+i*4.05; box(s,x,4.35,3.85,1.75,fill=LG if i!=1 else TINT)
    tb(s,x+0.18,4.47,3.5,0.3,a,size=12.5,bold=True,font=SERIF); tb(s,x+0.18,4.82,3.5,1.2,b2,size=10.5,color=GREY,ls=1.0)

# ---------- 9 RESULTS
s=SL[8]; title(s,'EXPECTED RESULTS & IMPACT'); head(s,'Early results from our own Qiskit runs and navigation simulation.')
cd=CategoryChartData(); keys=sorted(D['counts']); cd.categories=keys; cd.add_series('Shots',[D['counts'][k] for k in keys])
gf=s.shapes.add_chart(CT.COLUMN_CLUSTERED,I(0.6),I(1.85),I(4.3),I(3.05),cd); ch=gf.chart; style_chart(ch,8)
ch.has_title=True; ch.chart_title.text_frame.text='Grover on Aer · 4,096 shots'; f=ch.chart_title.text_frame.paragraphs[0].runs[0].font; f.size=Pt(11); f.bold=True; f.color.rgb=INK
ch.has_legend=False; pl=ch.plots[0]; pl.gap_width=40
ser=pl.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb=C(0xD9,0xD9,0xDF)
for i,k in enumerate(keys):
    if k in ('0110','1011'): pt=ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb=PINK
ch.value_axis.has_major_gridlines=False; ch.value_axis.visible=False; ch.category_axis.tick_labels.font.size=Pt(7); ch.category_axis.format.line.color.rgb=MID
tb(s,0.7,4.95,4.2,0.6,[('0110 + 1011 = 94% of shots. ',{'bold':True,'color':DEEP}),('Grover returns 2 look-alike cells; QSVM breaks the tie.',{})],size=10.5,color=GREY)
cd=CategoryChartData(); cd.categories=[f'{t}s' for t in D['t']]
cd.add_series('Inertial only',D['ins']); cd.add_series('VOID-NAV · MEMS magnetometer',D['mems']); cd.add_series('VOID-NAV · simulated NV sensor',D['nv'])
gf=s.shapes.add_chart(CT.LINE,I(5.05),I(1.85),I(4.35),I(3.05),cd); ch=gf.chart; style_chart(ch,8)
ch.has_title=True; ch.chart_title.text_frame.text='Error after GPS loss (m, log scale)'; f=ch.chart_title.text_frame.paragraphs[0].runs[0].font; f.size=Pt(11); f.bold=True; f.color.rgb=INK
ch.has_legend=True; ch.legend.position=LP.BOTTOM; ch.legend.include_in_layout=False; ch.legend.font.size=Pt(8)
for sr,col in zip(ch.plots[0].series,[MID,PINK,INK]): sr.format.line.color.rgb=col; sr.format.line.width=Pt(2.25); sr.smooth=True
va=ch.value_axis; va._element.get_or_add_scaling().get_or_add_logBase().set('val','10') if False else None
sc=va._element.find(qn('c:scaling')); lb=sc.makeelement(qn('c:logBase'),{'val':'10'}); sc.insert(0,lb); va.minimum_scale=0.1; va.maximum_scale=1000
va.major_gridlines.format.line.color.rgb=C(0xE6,0xE6,0xEA); ch.category_axis.tick_labels.font.size=Pt(7)
img(s,'kernel.png',9.75,1.9,h=2.6)
tb(s,9.6,4.55,3.1,0.95,[('Quantum kernel matrix. ',{'bold':True,'color':DEEP}),('Same-zone fingerprints overlap ~2.8× more than cross-zone ones (0.24 vs 0.09).',{})],size=10.5,color=GREY)
kp=[('~50×','lower drift vs inertial-only (MEMS model)'),('~6×','further gain with a quantum (NV) magnetometer'),('0 W','radio emission: fully passive')]
for i,(k,v) in enumerate(kp):
    x=0.7+i*4.05; box(s,x,5.5,3.85,0.62,fill=TINT)
    tb(s,x+0.15,5.5,1.1,0.62,k,size=20,bold=True,font=SERIF,color=PINK,anchor=V.MIDDLE); tb(s,x+1.25,5.5,2.5,0.62,v,size=10,color=INK,anchor=V.MIDDLE)
# ---------- 10 FUTURE
s=SL[9]; title(s,'FUTURE SCOPE & CONCLUSION'); head(s,'From a hackathon payload to a field-ready quantum PNT module.')
rm=[('Now · TRL 3','Hackathon','Payload prototype, 16-cell map, 4 quantum blocks on Aer + one IBM hardware run, live dashboard.'),
('6 months · TRL 4','Scale the map','Campus and outdoor anomaly maps (NGRI / EMAG2), larger Grover circuits, tethered drone trials.'),
('18 months · TRL 5','Quantum sensor','Swap in an NV-diamond magnetometer with an NQM sensing-hub partner; quantum IMU fusion.'),
('36 months · TRL 6','Defence trials','Ruggedised module with DRDO / iDEX partners on UAVs, loitering systems and vehicles.')]
c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,I(0.95),I(2.2),I(12.4),I(2.2)); c.line.color.rgb=INK; c.line.width=Pt(1.25)
for i,(a,b2,t2) in enumerate(rm):
    x=0.7+i*3.05; circ(s,x+0.1,2.03,0.34,'',fill=PINK if i==0 else WHITE)
    if i: s.shapes[-1].line.color.rgb=PINK; s.shapes[-1].line.width=Pt(2)
    eyebrow(s,x,2.55,a,w=2.9); tb(s,x,2.82,2.85,0.35,b2,size=14,bold=True,font=SERIF); tb(s,x,3.2,2.85,1.0,t2,size=10.5,color=GREY,ls=1.0)
box(s,0.7,4.45,12.0,1.65,fill=INK)
tb(s,1.0,4.6,8.2,0.9,'"When GPS goes dark, the ground still knows where you are."',size=22,font=SERIF,italic=True,color=WHITE)
tb(s,1.0,5.45,8.2,0.5,'VOID-NAV: quantum sensing, quantum search, quantum learning and quantum planning, flying on one passive payload.',size=11.5,color=C(0xD6,0xD6,0xDE))
tb(s,9.5,4.7,3.0,0.5,'Team VOID',size=24,font=SERIF,bold=True,color=WHITE,align=A.RIGHT)
box(s,9.45,5.52,3.05,0.4,fill=WHITE,shape=S.ROUNDED_RECTANGLE,radius=0.3)
tb(s,9.45,5.52,3.05,0.4,[('Video: ',{'color':PINK}),('https://youtu.be/BV2KYV-yhL0',{'link':'https://youtu.be/BV2KYV-yhL0','color':DEEP})],size=11,align=A.CENTER,anchor=V.MIDDLE,bold=True)
tb(s,9.5,5.2,3.0,0.3,'THANK YOU · QUESTIONS WELCOME',size=9,font=MONO,color=C(0xF9,0xA8,0xD8),align=A.RIGHT,spc=150)
p.save('/home/user/Quantique/deck/VOID-NAV_QiskitFallFest.pptx'); print('saved')
