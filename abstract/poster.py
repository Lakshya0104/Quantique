from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.shapes import MSO_SHAPE as S, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN as A, MSO_ANCHOR as V
from pptx.oxml.ns import qn
NAVY=C(0x1B,0x3C,0x74); PUR=C(0x6C,0x1C,0xF0); INK=C(0x1A,0x1A,0x1A); GREY=C(0x55,0x5B,0x66); TINT=C(0xF1,0xEC,0xFE); LG=C(0xF4,0xF5,0xF8); WHITE=C(255,255,255); PINK=C(0xF2,0x37,0xA6)
p=Presentation('a.pptx'); s=list(p.slides)[0]
kill=['TextBox 51','TextBox 52','TextBox 53','TextBox 54','TextBox 55','TextBox 56','TextBox 57','TextBox 58','TextBox 59','TextBox 60','Group 15']
for sh in list(s.shapes):
    if sh.name in kill: sh._element.getparent().remove(sh._element)
def settext(name,t,size=None):
    def walk(shs):
        for sh in shs:
            if sh.shape_type==6: walk(sh.shapes)
            elif sh.name==name:
                pp=sh.text_frame.paragraphs[0]; rs=pp.runs; rs[0].text=t
                for r in rs[1:]: r._r.getparent().remove(r._r)
                if size: rs[0].font.size=Pt(size)
    walk(s.shapes)
settext('TextBox 49','VOID-NAV: Quantum-Assisted Navigation for GPS-Denied Drones',33)
settext('TextBox 36','TEAM VOID',14)
def tb(x,y,w,h,paras,size=15,color=INK,font='Calibri',align=A.LEFT,anchor=V.TOP,after=4,bold=False):
    b=s.shapes.add_textbox(I(x),I(y),I(w),I(h)); tf=b.text_frame; tf.word_wrap=True
    for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,m,0)
    tf.vertical_anchor=anchor
    if isinstance(paras,str): paras=[paras]
    for i,para in enumerate(paras):
        pp=tf.paragraphs[0] if i==0 else tf.add_paragraph(); pp.alignment=align; pp.space_after=Pt(after)
        if isinstance(para,str): para=[(para,{})]
        for t,o in para:
            r=pp.add_run(); r.text=t; f=r.font; f.name=o.get('font',font); f.size=Pt(o.get('size',size)); f.bold=o.get('bold',bold); f.italic=o.get('italic',False); f.color.rgb=o.get('color',color)
            if o.get('link'): r.hyperlink.address=o['link']
    return b
def box(x,y,w,h,fill=None,line=None,shape=S.ROUNDED_RECTANGLE,lw=1,rad=0.08):
    r=s.shapes.add_shape(shape,I(x),I(y),I(w),I(h))
    if fill is None: r.fill.background()
    else: r.fill.solid(); r.fill.fore_color.rgb=fill
    if line is None: r.line.fill.background()
    else: r.line.color.rgb=line; r.line.width=Pt(lw)
    r.shadow.inherit=False
    if shape==S.ROUNDED_RECTANGLE: r.adjustments[0]=rad
    return r
def arrow(x1,y1,x2,y2,col=PUR):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,I(x1),I(y1),I(x2),I(y2)); c.line.color.rgb=col; c.line.width=Pt(2)
    ln=c.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}))
def pic(path,x,y,w=None,h=None,maxw=None,maxh=None):
    im=s.shapes.add_picture(path,I(x),I(y)); r=im.width/im.height
    W,H=I(maxw),I(maxh)
    if W/H>r: im.height=H; im.width=int(H*r)
    else: im.width=W; im.height=int(W/r)
    im.left=int(I(x)+(W-im.width)/2); im.top=int(I(y)+(H-im.height)/2); return im
B=lambda t:(t,{'bold':True})
# ===== A: problem / solution
tb(1.11,8.95,10.3,2.6,[
 [('When GPS is jammed, blocked by terrain or simply unavailable, a drone falls back on its inertial sensors (IMU). Their small errors add up, and the position ',{}),B('drifts without bound'),('. Nothing on board can tell by how much.',{})],
 [B('Who is affected: '),('defence UAVs, loitering systems, border and coastal patrol, and any drone flying in tunnels, forests or contested airspace.',{})],
 [B('Why current approaches fall short: '),('low-cost MEMS IMUs drift hundreds of metres within minutes; navigation-grade INS and cold-atom sensors cost lakhs to crores; radar or camera fixes emit signals or fail in darkness and smoke.',{})]],size=15,after=8)
stats=[('10⁻¹⁶ W','GPS power at the receiver: easy to lose'),('t²','inertial error grows with time squared'),('UC-085','quantum sensing for GPS-denied navigation')]
for i,(k,v) in enumerate(stats):
    x=1.11+i*3.5; box(x,11.85,3.3,2.3,fill=TINT)
    tb(x+0.2,12.0,2.9,0.9,k,size=30,color=PUR,font='Arial',bold=True)
    tb(x+0.2,12.95,2.9,1.1,v,size=13,color=GREY)
tb(12.81,8.95,10.3,1.4,[[B('VOID-NAV'),(' is a small, passive navigation payload that mounts on a drone. It reads the Earth\'s ',{}),B('magnetic fingerprint'),(', matches it to a pre-built magnetic map with ',{}),B('quantum algorithms'),(', and corrects inertial drift. No satellites, no emissions.',{})]],size=15)
pic('drone.png',12.7,10.35,maxw=5.6,maxh=3.8)
steps=[('1 · Map','Survey once with GPS on: fingerprint [Bx, By, Bz, |B|, ∇B] per 2 m cell'),('2 · Sense','GPS off: magnetometer, IMU, optical flow, barometer'),('3 · Match','Quantum engine finds the matching map cell'),('4 · Correct','Particle filter resets drift, sends position to autopilot')]
for i,(a,b) in enumerate(steps):
    y=10.45+i*0.93; tb(18.45,y,4.9,0.3,a,size=14,color=NAVY,font='Arial',bold=True); tb(18.45,y+0.33,4.9,0.6,b,size=12.5,color=GREY,after=0)
# ===== B: key features / how it works
feats=[('Full quantum pipeline','sensing, search, learning and planning, each with a distinct job'),('5-D magnetic fingerprint','matches field pattern and gradient, not a single value'),('Grover + QSVM cross-check','a correction is accepted only when both agree'),('Quantum-sensor ready','drop-in slot for an NV-diamond magnetometer'),('Passive and retrofittable','emits nothing; talks MAVLink to ArduPilot/PX4'),('Low cost','prototype under ₹4,000 in parts')]
for i,(a,b) in enumerate(feats):
    x=0.99+(i%2)*5.2; y=15.95+(i//2)*1.25
    box(x,y+0.08,0.28,0.28,fill=PUR,shape=S.OVAL)
    tb(x+0.45,y,4.6,0.32,a,size=14,color=INK,font='Arial',bold=True); tb(x+0.45,y+0.36,4.6,0.8,b,size=12.5,color=GREY,after=0)
flow=[('SENSE','50 Hz sensors'),('PRE-PROCESS','calibrate + fingerprint'),('QUANTUM','IPE · Grover · QSVM · QAOA'),('FUSE','particle filter'),('ACT','autopilot + dashboard')]
for i,(a,b) in enumerate(flow):
    x=13.12+i*2.08; bx=box(x,16.3,1.8,1.45,fill=PUR if i==2 else LG)
    tb(x+0.08,16.42,1.64,0.3,a,size=11.5,font='Arial',bold=True,color=WHITE if i==2 else NAVY,align=A.CENTER)
    tb(x+0.08,16.78,1.64,0.9,b,size=12,color=WHITE if i==2 else GREY,align=A.CENTER,after=0)
    if i<4: arrow(x+1.82,17.02,x+2.06,17.02)
tb(13.12,18.0,10.3,1.8,[[('Sensor data streams in at 50 Hz. Every second the quantum engine takes a window of readings and returns candidate map cells; the particle filter accepts a correction only when ',{}),B('Grover and the QSVM agree'),(', then sends position + confidence to the autopilot. Quantum runs alongside flight control, never blocking it.',{})]],size=13.5,color=INK)
# ===== C: quantum simulation (left) + circuits (right)
tb(0.99,20.7,10.3,1.3,[[('The full quantum pipeline runs in a Qiskit notebook on the ',{}),B('Qiskit Aer'),(' simulator, on a 4×4 magnetic map. The simulation is part of the navigation loop: its outputs directly correct the drone\'s position.',{})]],size=14)
rows=[('Simulated block','What it computes','Verified result'),('Ramsey + IPE','field phase from an NV-diamond qubit','16/16 map cells recovered exactly (6-bit)'),('Grover (4 qubits)','matching map cell','96.1% of shots on correct cell 1011'),('QSVM (ZZ kernel)','zone: corridor / stairwell / lab','99.0% test accuracy'),('QAOA (p = 1)','best waypoint pair','matches brute-force optimum'),('Navigation loop','drift with vs without quantum fixes','mean error 2.10 m → 0.40 m (5.3×)')]
t=s.shapes.add_table(len(rows),3,I(0.99),I(22.05),I(10.3),I(3.2)).table
for j,wd in enumerate([2.6,3.7,4.0]): t.columns[j].width=I(wd)
for i,r in enumerate(rows):
    for j,v in enumerate(r):
        c=t.cell(i,j); c.text=v; c.margin_left=c.margin_right=I(0.08); c.margin_top=c.margin_bottom=I(0.05)
        c.fill.solid(); c.fill.fore_color.rgb=NAVY if i==0 else (WHITE if i%2 else LG)
        f=c.text_frame.paragraphs[0].runs[0].font; f.size=Pt(12); f.name='Calibri'; f.bold=(i==0 or j==2); f.color.rgb=WHITE if i==0 else (PUR if j==2 else INK)
pic('/tmp/claude-0/abs/nb4.png',0.99,25.4,maxw=10.3,maxh=2.95)
tb(0.99,28.4,10.3,0.3,'Navigation loop output from the notebook: inertial-only drift (grey) vs VOID-NAV with Grover + QSVM fixes (pink).',size=11,color=GREY)
circ=[('Q1 · Quantum sensing: Ramsey + Iterative Phase Estimation','ipe.png','H creates superposition; the field imprints phase φ = 2πγBτ; phase kickback + IPE reads φ bit by bit.'),
('Q2 · Grover\'s search: oracle + diffuser, 4 qubits','grover.png','Oracle U_f flips matching cells; diffuser 2|s⟩⟨s|−I amplifies them; ~√N queries.'),
('Q3 · Quantum-kernel SVM: ZZ feature map','zz.png','K(x,x′) = |⟨φ(x)|φ(x′)⟩|² from statevectors, fed to a classical SVM.'),
('Q4 · QAOA route planning, p = 1','qaoa.png','Cost Hamiltonian from a QUBO; mixer ΣXᵢ; (γ, β) optimised on Aer.')]
for i,(a,im,b) in enumerate(circ):
    y=21.05+i*1.9
    tb(13.06,y,10.3,0.3,a,size=13,color=NAVY,font='Arial',bold=True)
    pic(im,13.06,y+0.35,maxw=5.4,maxh=1.35)
    tb(18.7,y+0.4,4.66,1.3,b,size=12,color=GREY,after=0,anchor=V.MIDDLE)
# ===== D: bottom row
tools=['Qiskit 2.x SDK (circuits, primitives)','Qiskit Aer (AerSimulator)','Qiskit quantum_info (Statevector kernels)','Qiskit circuit library (zz_feature_map, DiagonalGate)','IBM Quantum Platform via qiskit-ibm-runtime (hardware run planned)','Also: scikit-learn, NumPy, Matplotlib, ArduPilot SITL']
tb(1.19,30.15,6.9,4.6,[[('• ',{'color':PUR,'bold':True}),(x,{})] for x in tools],size=16,after=11)
imp=['GPS-independent, jam-proof positioning for defence drones','Bounded error instead of unbounded drift (5.3× lower in simulation)','Passive: no emissions to detect','Low-cost payload that retrofits existing drones','Dual use: disaster response, mining, tunnels, maritime']
tb(8.17,30.15,6.9,4.6,[[('• ',{'color':PUR,'bold':True}),(x,{})] for x in imp],size=16,after=11)
fut=['Run Grover and IPE on IBM Quantum hardware','Replace the synthetic map with our real magnetometer survey','Larger maps: more qubits, outdoor anomaly data','Swap in a real NV-diamond quantum magnetometer','Field trials with defence partners (DRDO / iDEX)']
tb(15.1,30.15,7.6,4.0,[[('• ',{'color':PUR,'bold':True}),(x,{})] for x in fut],size=16,after=11)
VID='https://youtu.be/BV2KYV-yhL0'
tb(1.19,34.55,15,0.4,[[B('Video: '),(VID,{'link':VID,'color':PUR}),('     Notebook: ',{'bold':True}),('VOID-NAV_Quantum_Simulation.ipynb',{'color':PUR})]],size=15)
p.save('/home/user/Quantique/abstract/VOID-NAV_QuantumXpo_Abstract.pptx'); print('ok')
