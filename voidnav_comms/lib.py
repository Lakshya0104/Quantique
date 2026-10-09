from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.shapes import MSO_SHAPE as S, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN as A, MSO_ANCHOR as V
from pptx.oxml.ns import qn
INK=C(0x16,0x19,0x1D); MUT=C(0x55,0x5D,0x66); LINE=C(0xC9,0xD0,0xD6); TEAL=C(0x1F,0x5F,0x6B); TL=C(0xE6,0xF0,0xF1)
ORANGE=C(0xE0,0x7B,0x1A); OL=C(0xFD,0xF1,0xE6); GREEN=C(0x2E,0x8B,0x57); GL=C(0xE7,0xF4,0xEC); RUST=C(0xB5,0x53,0x2A); RL=C(0xFB,0xEE,0xE8)
GREY=C(0xF3,0xF5,0xF6); WHITE=C(255,255,255); DARK=C(0x24,0x2A,0x30); SLATE=C(0x8A,0x93,0x9B)
SERIF='Cambria'; SANS='Calibri'; MONO='Consolas'
def tb(s,x,y,w,h,paras,size=10,color=INK,font=SANS,bold=False,align=A.LEFT,anchor=V.TOP,after=0,ls=None,italic=False):
    b=s.shapes.add_textbox(I(x),I(y),I(w),I(h)); tf=b.text_frame; tf.word_wrap=True
    for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,m,0)
    tf.vertical_anchor=anchor
    if isinstance(paras,str): paras=[paras]
    for i,para in enumerate(paras):
        pp=tf.paragraphs[0] if i==0 else tf.add_paragraph(); pp.alignment=align; pp.space_after=Pt(after)
        if ls: pp.line_spacing=ls
        if isinstance(para,str): para=[(para,{})]
        elif isinstance(para,tuple): para=[para]
        for t,o in para:
            r=pp.add_run(); r.text=t; f=r.font; f.name=o.get('font',font); f.size=Pt(o.get('size',size)); f.bold=o.get('bold',bold); f.italic=o.get('italic',italic); f.color.rgb=o.get('color',color)
            if o.get('link'): r.hyperlink.address=o['link']
            if o.get('spc'): r._r.get_or_add_rPr().set('spc',str(o['spc']))
    return b
def B(t,**k): k.setdefault('bold',True); return (t,k)
def T(t,**k): return (t,k)
def box(s,x,y,w,h,fill=None,line=None,dash=False,lw=1,shape=S.RECTANGLE,rad=0.1,text=None,size=9,tcolor=INK,bold=False,align=A.CENTER,anchor=V.MIDDLE,font=SANS):
    r=s.shapes.add_shape(shape,I(x),I(y),I(w),I(h))
    if fill is None: r.fill.background()
    else: r.fill.solid(); r.fill.fore_color.rgb=fill
    if line is None: r.line.fill.background()
    else:
        r.line.color.rgb=line; r.line.width=Pt(lw)
        if dash: r.line.dash_style=7 if dash is True else dash
    r.shadow.inherit=False
    if shape==S.ROUNDED_RECTANGLE: r.adjustments[0]=rad
    tf=r.text_frame
    for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,m,I(0.03))
    tf.word_wrap=True; tf.vertical_anchor=anchor
    if text is not None:
        lines=text if isinstance(text,list) else [text]
        for i,l in enumerate(lines):
            pp=tf.paragraphs[0] if i==0 else tf.add_paragraph(); pp.alignment=align
            if isinstance(l,str): l=[(l,{})]
            elif isinstance(l,tuple): l=[l]
            for t,o in l:
                rr=pp.add_run(); rr.text=t; f=rr.font; f.size=Pt(o.get('size',size)); f.name=o.get('font',font); f.bold=o.get('bold',bold); f.color.rgb=o.get('color',tcolor); f.italic=o.get('italic',False)
    return r
def line(s,x1,y1,x2,y2,color=INK,w=1.25,dash=None,head=True,tail=False):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,I(x1),I(y1),I(x2),I(y2)); c.line.color.rgb=color; c.line.width=Pt(w)
    if dash: c.line.dash_style=dash
    ln=c.line._get_or_add_ln()
    if head: ln.append(ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}))
    if tail: ln.append(ln.makeelement(qn('a:headEnd'),{'type':'triangle','w':'med','len':'med'}))
    return c
def poly(s,pts,color=INK,w=1.25,dash=None,head=True):
    for i in range(len(pts)-1):
        line(s,*pts[i],*pts[i+1],color=color,w=w,dash=dash,head=(head and i==len(pts)-2))
def circ(s,x,y,d,txt='',fill=TEAL,fc=WHITE,size=9,line_c=None,bold=True):
    c=box(s,x,y,d,d,fill=fill,line=line_c,shape=S.OVAL,text=txt,size=size,tcolor=fc,bold=bold)
    c.text_frame.margin_left=c.text_frame.margin_right=0; return c
def eyebrow(s,x,y,t,color=TEAL,w=5,size=8.5): return tb(s,x,y,w,0.2,[T(t.upper(),spc=180,bold=True)],size=size,color=color,font=MONO)
def header(s,num,title,sub):
    eyebrow(s,0.45,0.28,f'{num} · VOID-NAV · Tech Horizon 2.0',w=8)
    tb(s,0.45,0.48,12.4,0.5,title,size=25,font=SERIF,bold=True)
    if sub: tb(s,0.45,0.98,12.4,0.3,sub,size=12,font=SERIF,italic=True,color=MUT)
def footer(s,n):
    box(s,0.45,7.17,12.43,0.012,fill=LINE)
    tb(s,0.45,7.22,6,0.2,[T('TEAM VOID · VOID-NAV · DISASTER-RESILIENT SOS COMMUNICATION',spc=120)],size=7,color=SLATE,font=MONO)
    tb(s,11.9,7.22,1,0.2,[T(f'{n:02d} / 06',spc=120)],size=7,color=SLATE,font=MONO,align=A.RIGHT)
def chip(s,x,y,t,fill,fc,w=None,size=7.5):
    w=w or 0.12+len(t)*0.062
    return box(s,x,y,w,0.2,fill=fill,shape=S.ROUNDED_RECTANGLE,rad=0.5,text=[T(t,bold=True)],size=size,tcolor=fc)
def person(s,x,y,col=RUST,sc=1.0):
    box(s,x+0.05*sc,y,0.1*sc,0.1*sc,fill=col,shape=S.OVAL)
    box(s,x,y+0.11*sc,0.2*sc,0.14*sc,fill=col,shape=S.ROUNDED_RECTANGLE,rad=0.45)
def node(s,x,y,label,fill=TEAL,w=0.62,h=0.36,sub=None,size=8):
    r=box(s,x,y,w,h,fill=fill,shape=S.ROUNDED_RECTANGLE,rad=0.15,text=[[B(label,size=size,color=WHITE)]]+([[T(sub,size=6.5,color=WHITE)]] if sub else []),size=size,tcolor=WHITE)
    return r
def relay(s,x,y,label,col=SLATE):
    # mast + antenna circle
    line(s,x+0.15,y+0.12,x+0.15,y+0.42,color=col,w=1.5,head=False)
    circ(s,x+0.04,y,0.22,'',fill=WHITE,line_c=col)
    circ(s,x+0.1,y+0.06,0.1,'',fill=col)
    tb(s,x-0.2,y+0.44,0.7,0.2,label,size=7.5,bold=True,color=INK,align=A.CENTER)
